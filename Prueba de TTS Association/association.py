import pandas as pd
import re
from openpyxl import load_workbook

def normalizar_texto(texto):
    """Normaliza texto: quita puntuación, mayúsculas y espacios extra"""
    if pd.isna(texto) or texto == '':
        return ""
    # Solo letras, números y espacios
    texto = re.sub(r'[^\w\s]', '', str(texto).lower())
    # Quitar espacios múltiples
    texto = re.sub(r'\s+', ' ', texto)
    return texto.strip()

def extraer_datos_txt(linea):
    """
    Extrae tiempo y diálogo del TXT
    Formato: [Speaker X] (MM:SS): Diálogo
    """
    patron = r'\[.*?\]\s*\((\d+:\d+)\):\s*(.+)'
    match = re.match(patron, linea)
    
    if match:
        tiempo = match.group(1)
        dialogo = match.group(2).strip()
        return tiempo, dialogo
    return None, None

def procesar_archivos(txt_path, excel_path, output_path):
    """
    Asocia tiempos del TXT al Excel usando las primeras 3 palabras
    """
    print("="*60)
    print("INICIANDO PROCESO DE ASOCIACIÓN")
    print("="*60)
    
    # 1. Leer archivo TXT
    print("\n1. Leyendo archivo TXT...")
    txt_dialogues = []
    
    with open(txt_path, 'r', encoding='utf-8') as f:
        for linea in f:
            linea = linea.strip()
            if linea:
                tiempo, dialogo = extraer_datos_txt(linea)
                if tiempo and dialogo:
                    dialogo_normalizado = normalizar_texto(dialogo)
                    txt_dialogues.append({
                        'tiempo': tiempo,
                        'dialogo_original': dialogo,
                        'dialogo_normalizado': dialogo_normalizado
                    })
    
    print(f"   ✓ Total diálogos en TXT: {len(txt_dialogues)}")
    
    # 2. Leer archivo Excel
    print("\n2. Leyendo archivo Excel...")
    workbook = load_workbook(excel_path)
    sheet = workbook.active
    
    # Buscar columnas
    original_dialogue_col = None
    timestamps_col = None
    
    for col in range(1, sheet.max_column + 1):
        cell_value = sheet.cell(row=1, column=col).value
        if cell_value == "DIÁLOGO ORIGINAL":
            original_dialogue_col = col
        elif cell_value == "TIEMPOS":
            timestamps_col = col
    
    if original_dialogue_col is None:
        print("   ✗ ERROR: No se encontró la columna 'DIÁLOGO ORIGINAL'")
        return
    
    if timestamps_col is None:
        print("   ✗ ERROR: No se encontró la columna 'TIEMPOS'")
        return
    
    print(f"   ✓ Columna 'DIÁLOGO ORIGINAL': {original_dialogue_col}")
    print(f"   ✓ Columna 'TIEMPOS': {timestamps_col}")
    
    # 3. Obtener diálogos del Excel
    print("\n3. Extrayendo diálogos del Excel...")
    excel_dialogues = []
    
    for row in range(2, sheet.max_row + 1):
        original_dialogue = sheet.cell(row=row, column=original_dialogue_col).value
        if original_dialogue:
            normalized_dialogue = normalizar_texto(original_dialogue)
            
            # Filtrar líneas muy cortas (1-2 palabras)
            words = normalized_dialogue.split()
            if len(words) <= 2:
                continue
            
            # Filtrar expresiones cortas comunes
            short_expressions = ['hmmm', 'mmmm', 'hmm', 'mmm', 'same', 'crunchy', 
                               'yum', 'yeah', 'ok', 'wow', 'oh', 'ah', 'um', 'uh']
            if normalized_dialogue in short_expressions:
                continue
            
            excel_dialogues.append({
                'row': row,
                'dialogo_original': original_dialogue,
                'dialogo_normalizado': normalized_dialogue
            })
    
    print(f"   ✓ Total diálogos válidos en Excel: {len(excel_dialogues)}")
    
    # 4. Asociación secuencial por primeras 3 palabras
    print("\n4. Asociando tiempos (comparando primeras 3 palabras)...")
    print("-"*60)
    
    matches_found = 0
    txt_index = 0
    excel_index = 0
    
    while excel_index < len(excel_dialogues) and txt_index < len(txt_dialogues):
        excel_data = excel_dialogues[excel_index]
        txt_data = txt_dialogues[txt_index]
        
        # Tomar las primeras 3 palabras de cada diálogo
        excel_words = excel_data['dialogo_normalizado'].split()[:3]
        txt_words = txt_data['dialogo_normalizado'].split()[:3]
        
        excel_start = ' '.join(excel_words)
        txt_start = ' '.join(txt_words)
        
        # Verificar coincidencia
        is_match = False
        
        # Coincidencia si el inicio del Excel está en el TXT o viceversa
        if (excel_start in txt_start and len(excel_words) >= 2) or \
           (txt_start in excel_start and len(txt_words) >= 2):
            if len(excel_words) >= 2 and len(txt_words) >= 2:
                is_match = True
        
        if is_match:
            # Asignar timestamp
            row_num = excel_data['row']
            sheet.cell(row=row_num, column=timestamps_col).value = txt_data['tiempo']
            matches_found += 1
            
            print(f"✓ Match #{matches_found}:")
            print(f"  Tiempo: {txt_data['tiempo']}")
            print(f"  Excel (fila {row_num}): {excel_data['dialogo_original'][:70]}...")
            print(f"  TXT: {txt_data['dialogo_original'][:70]}...")
            print(f"  Inicio Excel: '{excel_start}'")
            print(f"  Inicio TXT: '{txt_start}'")
            print()
            
            # Avanzar ambos índices
            txt_index += 1
            excel_index += 1
        else:
            # No hay coincidencia, avanzar solo el Excel
            excel_index += 1
        
        # Si terminamos el Excel sin suficientes coincidencias, avanzar TXT
        if excel_index >= len(excel_dialogues) and txt_index < len(txt_dialogues) - 1:
            txt_index += 1
            excel_index = 0
    
    print("-"*60)
    print(f"\n5. Resumen:")
    print(f"   ✓ Coincidencias encontradas: {matches_found}")
    print(f"   ✓ Total filas Excel procesadas: {len(excel_dialogues)}")
    print(f"   ✓ Porcentaje de éxito: {(matches_found/len(excel_dialogues)*100):.1f}%")
    
    # 6. Guardar resultado
    print(f"\n6. Guardando resultado en: {output_path}")
    workbook.save(output_path)
    print("   ✓ Archivo guardado exitosamente")
    
    print("\n" + "="*60)
    print("PROCESO COMPLETADO")
    print("="*60)

# Ejecutar
if __name__ == "__main__":
    # AJUSTA ESTAS RUTAS
    txt_path = "diarization_GEMINI.txt"
    excel_path = "dialogos_unificados.xlsx"
    output_path = "AW 301_con_tiempos.xlsx"
    
    procesar_archivos(txt_path, excel_path, output_path)